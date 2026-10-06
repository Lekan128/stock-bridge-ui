"""
Builds the two free downloads (LANDING_PAGE_PLAN.md §5, "the lead magnet that is also the funnel"):

  public/downloads/procurepaddy-inventory-template.xlsx
      The API's own import template (import-template.xlsx: GET /api/products/template, saved here),
      with a "Read me" tab in front and two working tabs behind Products: a stock in-and-out log
      and a count sheet. The Products tab is untouched and stays the first tab that isn't help
      ("Read me" is one of the help names the importer steps over: TemplateConventions), so the
      file a shop outgrows is the file Procurepaddy imports.
  public/downloads/procurepaddy-stock-count-sheet.xlsx
      A printable count sheet: product, expected, counted, difference, who counted.

Run after the import template changes:  python3 scripts/lead-magnets/build.py   (needs openpyxl)
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
OUT = HERE.parent.parent / "public" / "downloads"
NAVY = "08205B"
HEADER = PatternFill("solid", fgColor=NAVY)
TINT = PatternFill("solid", fgColor="EFF5FF")
THIN = Side(style="thin", color="D4D4D8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def header_row(ws, headers, widths):
    for column, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=column, value=title)
        cell.fill = HEADER
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def start_here(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 96
    lines = [
        ("Free inventory template, from Procurepaddy", Font(bold=True, size=18, color=NAVY)),
        ("Keep your shop's stock in Excel, properly.", Font(size=12, color="3F3F46")),
        ("", None),
        ("The tabs", Font(bold=True, size=13, color=NAVY)),
        ("1.  Products: one row per product. How it comes (bag, carton), what is inside one, and how many you have.", None),
        ("2.  Stock in and out: every delivery and every sale or issue, one line each, with who recorded it.", None),
        ("3.  Count sheet: print it, count a shelf, and see the difference from your records.", None),
        ("4.  How to fill this in: the notes for the Products tab, column by column.", None),
        ("", None),
        ("Three rules that keep the numbers right", Font(bold=True, size=13, color=NAVY)),
        ("•  Write every movement the same day, with a name beside it. A record without a name can't be checked.", None),
        ("•  Count one shelf every week, not the whole shop twice a year. Small differences are easy to explain.", None),
        ("•  Keep one copy of this file. Two copies means two different answers.", None),
        ("", None),
        ("When this gets too much", Font(bold=True, size=13, color=NAVY)),
        ("This file is the format Procurepaddy imports. When your staff need to record on their own phones, at the same", None),
        ("time, with no network, upload this file to Procurepaddy and your products are in. Every change then carries a name", None),
        ("and a time, and a count shows the difference straight away. Free to start: procurepaddy.com", Font(bold=True, color=NAVY)),
        ("WhatsApp +234 818 410 3312, 8am to 6pm, Monday to Saturday.", None),
    ]
    for row, (text, font) in enumerate(lines, start=2):
        cell = ws.cell(row=row, column=2, value=text)
        if font:
            cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def movements(ws):
    header_row(
        ws,
        ["Date", "Product name", "In or out", "How many", "Counted in (bag, carton, kg…)", "Price for one (₦)", "Supplier or customer", "Recorded by", "Note"],
        [12, 32, 11, 11, 24, 16, 24, 16, 30],
    )
    in_out = DataValidation(type="list", formula1='"In,Out"', allow_blank=True)
    ws.add_data_validation(in_out)
    in_out.add("C2:C5000")
    examples = [
        ("2026-10-05", "Rice (Mama Gold)", "In", 10, "Bag", 42000, "Mama Gold depot", "Chidi", "Delivery note 0143"),
        ("2026-10-05", "Rice (Mama Gold)", "Out", 2, "Bag", 45000, "Mrs Bello", "Amaka", ""),
    ]
    for row, values in enumerate(examples, start=2):
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.font = Font(italic=True, color="71717A")
    ws.cell(row=4, column=1, value="The two grey rows are examples. Type over them.").font = Font(italic=True, color="71717A")


def count_sheet(ws, title_rows=True):
    ws.sheet_view.showGridLines = False
    start = 1
    if title_rows:
        ws.cell(row=1, column=1, value="Stock count").font = Font(bold=True, size=16, color=NAVY)
        ws.cell(row=2, column=1, value="Shop: ____________________    Shelf or area: ____________________    Date: ____________").font = Font(color="3F3F46")
        ws.cell(row=3, column=1, value="Counted by: ____________________    Checked by: ____________________").font = Font(color="3F3F46")
        start = 5
    headers = ["Product", "Counted in", "Should be (from records)", "Counted", "Difference", "Note"]
    widths = [34, 14, 16, 12, 12, 30]
    for column, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=start, column=column, value=title)
        cell.fill = HEADER
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[start].height = 32
    for row in range(start + 1, start + 41):
        for column in range(1, 7):
            cell = ws.cell(row=row, column=column)
            cell.border = BOX
            if row % 2 == 0:
                cell.fill = TINT
        diff = ws.cell(row=row, column=5, value=f'=IF(OR(C{row}="",D{row}=""),"",D{row}-C{row})')
        diff.alignment = Alignment(horizontal="right")
        ws.row_dimensions[row].height = 22
    red = PatternFill("solid", fgColor="FDE2E2")
    ws.conditional_formatting.add(f"E{start + 1}:E{start + 40}", CellIsRule(operator="lessThan", formula=["0"], fill=red))
    ws.print_title_rows = f"{start}:{start}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Free from procurepaddy.com · Count one shelf a week, not the whole shop twice a year"
    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def inventory_template():
    wb = load_workbook(HERE / "import-template.xlsx")
    wb.properties.creator = "Procurepaddy"
    wb.properties.title = "Free inventory template"
    start = wb.create_sheet("Read me", 0)
    start_here(start)
    after_products = wb.sheetnames.index("Products") + 1
    movement = wb.create_sheet("Stock in and out", after_products)
    movements(movement)
    count = wb.create_sheet("Count sheet", after_products + 1)
    count_sheet(count)
    wb["_lookups"].sheet_state = "hidden"
    wb.active = 0
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet.title == "Read me"
    wb.save(OUT / "procurepaddy-inventory-template.xlsx")


def stock_count_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "Count sheet"
    count_sheet(ws)
    wb.properties.creator = "Procurepaddy"
    wb.properties.title = "Stock count sheet"
    wb.save(OUT / "procurepaddy-stock-count-sheet.xlsx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    inventory_template()
    stock_count_sheet()
    for file in sorted(OUT.glob("*.xlsx")):
        print(f"{file.name}: {file.stat().st_size:,} bytes")
